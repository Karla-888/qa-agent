import sys, pathlib, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

RAIZ = pathlib.Path(__file__).parent.parent

META = {
    "proyecto": "qa-agent - Generacion asistida de casos de prueba",
    "sistema": "SauceDemo (sitio de practica)",
    "modulo": "Autenticacion / Login",
    "area": "Portafolio personal",
    "fecha": "2026-10-14",
    "objetivo": "Validar el funcionamiento del mecanismo de autenticacion, cubriendo flujos basicos, alternos y de excepcion.",
    "componentes": "Modulo de Login. Campos usuario y contrasena, boton Iniciar sesion, mensajeria de error.",
}
PRECONDICIONES_GRALES = [
    "El sistema se encuentra desplegado y accesible en el entorno de pruebas.",
    "Existen usuarios de prueba previamente configurados.",
    "Se cuenta con conectividad a los servicios de autenticacion.",
    "La pantalla de inicio de sesion esta implementada y accesible.",
]

SECCIONES = [
    ("Basico", "Flujos basicos", "Funcionalidad o actividad basica que el modulo debe realizar. Recorre el camino principal sin desviaciones."),
    ("Alterno", "Flujos alternos o de excepcion", "Escenario distinto al principal, resultado de una condicion alterna o de un error controlado."),
    ("Integridad", "Pruebas de integridad de datos y base de datos", "Pruebas de los metodos y procesos que acceden a la base de datos para verificar que se ejecutan correctamente."),
    ("Otras", "Otras pruebas", "Pruebas que no se ubican en las clasificaciones anteriores."),
]

COLS = ["No. CP", "Descripcion de la prueba", "Precondiciones del CP", "Pasos a seguir",
        "Datos de prueba", "Resultados esperados", "Estatus de CP", "Fecha de ejecucion",
        "Resultado del CP", "Hallazgos (Mantis)", "Comentarios"]

AZUL = "1F4E79"; GRIS = "D9D9D9"; CLARO = "F2F2F2"
BORDE = Border(*[Side(style="thin", color="999999")] * 4)


def est(c, bold=False, size=10, fill=None, color="000000", wrap=True, center=False):
    c.font = Font(name="Arial", size=size, bold=bold, color=color)
    c.alignment = Alignment(wrap_text=wrap, vertical="top",
                            horizontal="center" if center else "left")
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.border = BORDE


def exportar(archivo_csv: str, archivo_diag: str = None):
    df = pd.read_csv(RAIZ / "outputs" / archivo_csv).fillna("")

    diag = {}
    ruta_diag = RAIZ / "outputs" / (archivo_diag or "diagnostico.csv")
    if ruta_diag.exists():
        dd = pd.read_csv(ruta_diag).fillna("")
        for _, r in dd.iterrows():
            diag[str(r.get("caso", "")).strip()] = r

    wb = Workbook()
    ws = wb.active
    ws.title = "Casos de Prueba"
    ws.sheet_view.showGridLines = False

    for col, w in zip("ABCDEFGHIJK", [10, 42, 38, 46, 26, 46, 14, 16, 14, 18, 26]):
        ws.column_dimensions[col].width = w

    f = 1
    ws.merge_cells(f"A{f}:K{f}")
    est(ws.cell(f, 1), bold=True, size=14, fill=AZUL, color="FFFFFF", center=True)
    ws.cell(f, 1).value = "Documento de Apoyo (DA) - Casos de Prueba de Integracion y Negocio"
    ws.row_dimensions[f].height = 26

    f += 1
    ws.merge_cells(f"A{f}:K{f}")
    est(ws.cell(f, 1), bold=True, size=9, fill=GRIS, center=True)
    ws.cell(f, 1).value = "DA-F05-CasPruebaIntrNeg"

    f += 2
    for etiqueta, clave in [("Proyecto", "proyecto"), ("Sistema", "sistema"),
                            ("Modulo", "modulo"), ("Area", "area"), ("Fecha", "fecha")]:
        est(ws.cell(f, 1), bold=True, fill=CLARO)
        ws.cell(f, 1).value = etiqueta
        ws.merge_cells(f"B{f}:K{f}")
        est(ws.cell(f, 2))
        ws.cell(f, 2).value = META[clave]
        f += 1

    f += 1
    est(ws.cell(f, 1), bold=True, fill=GRIS)
    ws.cell(f, 1).value = "Objetivo"
    ws.merge_cells(f"B{f}:K{f}")
    est(ws.cell(f, 2))
    ws.cell(f, 2).value = META["objetivo"]
    ws.row_dimensions[f].height = 30
    f += 1
    est(ws.cell(f, 1), bold=True, fill=GRIS)
    ws.cell(f, 1).value = "Componentes"
    ws.merge_cells(f"B{f}:K{f}")
    est(ws.cell(f, 2))
    ws.cell(f, 2).value = META["componentes"]
    ws.row_dimensions[f].height = 30

    f += 2
    ws.merge_cells(f"A{f}:K{f}")
    est(ws.cell(f, 1), bold=True, fill=AZUL, color="FFFFFF")
    ws.cell(f, 1).value = "Precondiciones generales"
    f += 1
    for p in PRECONDICIONES_GRALES:
        ws.merge_cells(f"A{f}:K{f}")
        est(ws.cell(f, 1))
        ws.cell(f, 1).value = p
        f += 1

    f += 1
    inicio_tabla = f
    for i, nombre in enumerate(COLS, start=1):
        est(ws.cell(f, i), bold=True, fill=AZUL, color="FFFFFF", center=True)
        ws.cell(f, i).value = nombre
    ws.row_dimensions[f].height = 30
    f += 1

    total = ejecutados = fallaron = 0
    col_flujo = "flujo" if "flujo" in df.columns else None

    for clave, titulo, desc in SECCIONES:
        sub = df[df[col_flujo].astype(str).str.strip() == clave] if col_flujo else df
        ws.merge_cells(f"A{f}:K{f}")
        est(ws.cell(f, 1), bold=True, fill=GRIS)
        ws.cell(f, 1).value = f"{titulo} - {desc}"
        ws.row_dimensions[f].height = 28
        f += 1

        if len(sub) == 0:
            ws.merge_cells(f"A{f}:K{f}")
            est(ws.cell(f, 1))
            ws.cell(f, 1).value = "No aplica para este alcance."
            f += 1
        for _, r in sub.iterrows():
            total += 1
            d = diag.get(str(r.get("id_unico", "")).strip(), {})
            estatus = "Ejecutado" if len(d) else ""
            resultado = ""
            if len(d):
                ejecutados += 1
                clasif = str(d.get("clasificacion", "")).strip()
                resultado = "No paso" if clasif == "BUG" else "Reejecutar"
                fallaron += resultado == "No paso"
            valores = [
                r.get("no_cp", ""), r.get("descripcion", ""), r.get("precondiciones", ""),
                r.get("pasos", ""), r.get("datos_prueba", ""), r.get("resultados_esperados", ""),
                estatus, "", resultado,
                d.get("ticket_titulo", "") if len(d) else "",
                d.get("diagnostico", "") if len(d) else r.get("comentarios", ""),
            ]
            for i, v in enumerate(valores, start=1):
                est(ws.cell(f, i), center=(i in (1, 7, 8, 9)))
                ws.cell(f, i).value = str(v) if v is not None else ""
            f += 1
        f += 1

    f += 1
    ws.merge_cells(f"A{f}:C{f}")
    est(ws.cell(f, 1), bold=True, fill=AZUL, color="FFFFFF")
    ws.cell(f, 1).value = "Estadisticas"
    f += 1
    for etiqueta, valor in [("Total de casos de prueba disenados", total),
                            ("Ejecutados", ejecutados),
                            ("No ejecutados", total - ejecutados),
                            ("No pasaron (defecto real)", fallaron),
                            ("Pendientes de reejecucion", ejecutados - fallaron)]:
        ws.merge_cells(f"A{f}:B{f}")
        est(ws.cell(f, 1), bold=True, fill=CLARO)
        ws.cell(f, 1).value = etiqueta
        est(ws.cell(f, 3), center=True)
        ws.cell(f, 3).value = valor
        f += 1

    ws.freeze_panes = ws.cell(inicio_tabla + 1, 1)

    salida = RAIZ / "outputs" / f"DA-F05_{pathlib.Path(archivo_csv).stem.replace('_casos','')}.xlsx"
    wb.save(salida)
    print(f"{total} casos -> {salida}")


if __name__ == "__main__":
    exportar(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
